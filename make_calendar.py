"""
Скачивает расписание группы с timetable.spbu.ru (ту же Excel-выгрузку,
что и кнопка на сайте) и собирает из него календарь schedule.ics.

Запуск:  python make_calendar.py
Проверка на готовом файле:  python make_calendar.py --file путь/к/файлу.xlsx
"""

import hashlib
import json
import re
import sys
import time
from datetime import date, datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent
CONFIG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))
OUT = ROOT / "schedule.ics"

EXCEL_URL = "https://timetable.spbu.ru/StudentGroupEvents/ExcelWeek"
MSK = timedelta(hours=3)  # в Москве нет перехода на летнее время

MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
    "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
    "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
}
HEADERS = {
    "time": ("время", "time"),
    "name": ("название", "subject", "name"),
    "place": ("места проведения", "locations", "location"),
    "teacher": ("преподаватели", "educators", "teachers"),
}


# ---------- скачивание ----------

class SiteUnavailable(Exception):
    """Сайт СПбГУ не ответил или ответил ошибкой сервера — временный сбой."""


def download_week(monday, pauses=(20, 60)):
    import requests

    last_error = None
    for attempt in range(len(pauses) + 1):
        if attempt:
            print(f"  повтор через {pauses[attempt - 1]} с ({last_error})")
            time.sleep(pauses[attempt - 1])
        try:
            response = requests.get(
                EXCEL_URL,
                params={"studentGroupId": CONFIG["group_id"], "weekMonday": monday.isoformat()},
                cookies={"_culture": "ru"},
                headers={"User-Agent": "Mozilla/5.0 (spbu-timetable-to-calendar)"},
                timeout=(15, 90),  # 15 с на подключение, 90 с на ответ
            )
        except (requests.Timeout, requests.ConnectionError) as error:
            last_error = type(error).__name__
            continue
        if response.status_code >= 500 or response.status_code == 429:
            last_error = f"HTTP {response.status_code}"
            continue
        response.raise_for_status()  # прочие ошибки (403, 404) — не временные
        if not response.content.startswith(b"PK"):  # xlsx — это zip-архив
            last_error = "сайт вернул не Excel"
            continue
        return response.content
    raise SiteUnavailable(f"неделя с {monday:%d.%m.%Y}: {last_error}")


# ---------- разбор Excel ----------

def parse_day(text, monday):
    text = str(text).lower()
    m = re.search(r"(\d{1,2})\s+([a-zа-яё]+)", text) or None
    if m and m.group(2) in MONTHS:
        day, month = int(m.group(1)), MONTHS[m.group(2)]
    else:
        m = re.search(r"([a-z]+)\s+(\d{1,2})", text)
        if not m or m.group(1) not in MONTHS:
            return None
        day, month = int(m.group(2)), MONTHS[m.group(1)]
    # неделя может переходить через Новый год
    for year in (monday.year, monday.year + 1, monday.year - 1):
        try:
            d = date(year, month, day)
        except ValueError:
            continue
        if abs((d - monday).days) <= 7:
            return d
    return None


def parse_sheet(data, monday):
    ws = openpyxl.load_workbook(BytesIO(data), data_only=True).worksheets[0]

    # ищем строку с заголовками («Время», «Название», ...)
    cols, header_row = None, None
    for r in range(1, min(ws.max_row, 15) + 1):
        found = {}
        for c in range(1, ws.max_column + 1):
            v = str(ws.cell(r, c).value or "").strip().lower()
            for key, names in HEADERS.items():
                if v in names:
                    found[key] = c
        if "time" in found and "name" in found:
            cols, header_row = found, r
            break
    if cols is None:
        return []  # пустая неделя / нет занятий

    rows, day = [], None
    for r in range(header_row + 1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a:
            day = parse_day(a, monday)
        t = str(ws.cell(r, cols["time"]).value or "")
        name = str(ws.cell(r, cols["name"]).value or "").strip()
        m = re.match(r"\s*(\d{1,2}:\d{2})\s*[–—-]\s*(\d{1,2}:\d{2})", t)
        if not (day and m and name):
            continue
        get = lambda k: str(ws.cell(r, cols[k]).value or "").strip() if k in cols else ""
        rows.append(dict(day=day, t1=m.group(1), t2=m.group(2), name=" ".join(name.split()),
                         place=get("place"), teacher=get("teacher")))
    return rows


# ---------- сборка календаря ----------

def fmt_place(p):
    if not p:
        return ""
    if "информационно-коммуникационных" in p.lower():
        return "Онлайн"
    addr, _, room = p.rpartition(",")
    room = room.strip()
    if addr and re.fullmatch(r"[\w./-]{1,8}", room) and re.search(r"\d", room):
        return f"{addr.strip()}, ауд. {room}"
    return p


def esc(s):
    return s.replace("\\", "\\\\").replace(";", "\\;").replace(",", "\\,").replace("\n", "\\n")


def fold(line):
    if len(line.encode("utf-8")) <= 75:
        return line
    parts, cur, size, limit = [], "", 0, 75
    for ch in line:
        n = len(ch.encode("utf-8"))
        if size + n > limit:
            parts.append(cur)
            cur, size, limit = "", 0, 74
        cur += ch
        size += n
    parts.append(cur)
    return "\r\n ".join(parts)


def utc(d, hm):
    h, m = map(int, hm.split(":"))
    return (datetime(d.year, d.month, d.day, h, m) - MSK).strftime("%Y%m%dT%H%M%SZ")


def build_events(rows):
    """Строки из Excel → список событий: (начало в UTC, [строки VEVENT])."""
    # параллельные подгруппы (один предмет в одно время) — одним событием
    events = {}
    for row in rows:
        key = (row["day"], row["t1"], row["t2"], row["name"])
        variants = events.setdefault(key, [])
        if row not in variants:
            variants.append(row)

    blocks = []
    for (d, t1, t2, name), v in events.items():
        subj, sep, kind = name.rpartition(", ")
        title = f"{subj} ({kind})" if sep else name
        if len(v) == 1:
            loc = fmt_place(v[0]["place"])
            desc = "\n".join(x for x in [
                kind.capitalize() if sep else "",
                f"Преподаватель: {v[0]['teacher']}" if v[0]["teacher"] else "",
                f"Место: {loc}" if loc else "",
            ] if x)
        else:
            loc = " / ".join(dict.fromkeys(fmt_place(x["place"]) for x in v))
            desc = (kind.capitalize() + "\n" if sep else "") + "Параллельные подгруппы:\n" + \
                "\n".join(f"• {x['teacher']} — {fmt_place(x['place'])}" for x in v)
        uid = hashlib.md5(f"{d}|{t1}|{name}".encode()).hexdigest() + "@spbu-timetable"
        start = utc(d, t1)
        lines = ["BEGIN:VEVENT", f"UID:{uid}", f"DTSTAMP:{start}",
                 f"DTSTART:{start}", f"DTEND:{utc(d, t2)}",
                 f"SUMMARY:{esc(title)}"]
        if loc:
            lines.append(f"LOCATION:{esc(loc)}")
        if desc:
            lines.append(f"DESCRIPTION:{esc(desc)}")
        lines.append("END:VEVENT")
        blocks.append((start, lines))
    return blocks


def old_blocks():
    """События из прежнего schedule.ics — чтобы не терять недели, которые не скачались."""
    if not OUT.exists():
        return []
    text = OUT.read_bytes().decode("utf-8").replace("\r\n", "\n")
    text = text.replace("\n ", "")  # склеиваем перенесённые длинные строки
    blocks, cur = [], None
    for line in text.split("\n"):
        if line == "BEGIN:VEVENT":
            cur = [line]
        elif cur is not None:
            cur.append(line)
            if line == "END:VEVENT":
                start = next(x[8:] for x in cur if x.startswith("DTSTART:"))
                blocks.append((start, cur))
                cur = None
    return blocks


def week_of(start):
    local = datetime.strptime(start, "%Y%m%dT%H%M%SZ") + MSK
    return local.date() - timedelta(days=local.weekday())


def assemble(blocks):
    lines = [
        "BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//spbu-timetable//RU",
        "CALSCALE:GREGORIAN", "METHOD:PUBLISH",
        f"X-WR-CALNAME:{esc(CONFIG.get('calendar_name', 'Учёба'))}",
        "X-WR-TIMEZONE:Europe/Moscow",
        "REFRESH-INTERVAL;VALUE=DURATION:PT3H", "X-PUBLISHED-TTL:PT3H",
    ]
    for _, block in sorted(blocks, key=lambda b: (b[0], b[1])):
        lines += block
    lines.append("END:VCALENDAR")
    return "\r\n".join(fold(x) for x in lines) + "\r\n"


# ---------- отметка об успешных проверках ----------

STATUS = ROOT / "status.json"
ALERT_AFTER = timedelta(hours=24)  # молчим о сбоях сайта, если они короче суток


def read_last_ok():
    try:
        return datetime.fromisoformat(json.loads(STATUS.read_text())["last_ok"])
    except (OSError, ValueError, KeyError):
        return None


def mark_ok(now):
    last_ok = read_last_ok()
    if last_ok is None or now - last_ok > timedelta(hours=12):  # не чаще двух раз в сутки
        STATUS.write_text(json.dumps({"last_ok": now.isoformat(timespec="minutes")}) + "\n")


def main():
    if "--file" in sys.argv:  # проверка на уже скачанном файле
        path = Path(sys.argv[sys.argv.index("--file") + 1])
        blocks = build_events(parse_sheet(path.read_bytes(), date(2026, 10, 12)))
        OUT.write_text(assemble(blocks), encoding="utf-8", newline="")
        print(f"Готово: {len(blocks)} событий → {OUT.name}")
        return

    now = datetime.now(timezone.utc)
    today = (now + MSK).date()
    this_monday = today - timedelta(days=today.weekday())
    weeks = range(-CONFIG.get("weeks_back", 4), CONFIG.get("weeks_ahead", 10) + 1)

    rows, failed, in_a_row = [], set(), 0
    for i in weeks:
        monday = this_monday + timedelta(weeks=i)
        if in_a_row >= 2:  # сайт лежит — не мучаем его остальными неделями
            failed.add(monday)
            continue
        try:
            data = download_week(monday)
        except SiteUnavailable as error:
            print(f"::warning::Сайт СПбГУ не ответил: {error}. Оставляю прежние данные за эту неделю.")
            failed.add(monday)
            in_a_row += 1
            continue
        in_a_row = 0
        week_rows = parse_sheet(data, monday)
        print(f"Неделя с {monday:%d.%m.%Y}: занятий {len(week_rows)}")
        rows += week_rows
        time.sleep(1)  # не нагружаем сайт СПбГУ

    blocks = build_events(rows)
    kept = [b for b in old_blocks() if week_of(b[0]) in failed]
    OUT.write_text(assemble(blocks + kept), encoding="utf-8", newline="")
    print(f"Готово: {len(blocks)} новых событий, {len(kept)} оставлено из прежней версии")

    if not failed:
        mark_ok(now)
        return
    last_ok = read_last_ok()
    if last_ok and now - last_ok > ALERT_AFTER:
        print(f"::error::Расписание не удаётся скачать полностью с {last_ok:%d.%m %H:%M} UTC — "
              "больше суток. Календарь показывает последнюю удачную версию.")
        sys.exit(1)
    print(f"Сбой временный: не скачалось недель — {len(failed)}. Попробую при следующем запуске.")


if __name__ == "__main__":
    main()
